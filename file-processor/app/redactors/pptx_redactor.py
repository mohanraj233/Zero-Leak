import io
from typing import Tuple, List, Dict, Any
from app.detectors.presidio_detector import analyze_text

try:
    import pptx
except ImportError:
    pptx = None

def redact_pptx(file_bytes: bytes, filename: str) -> Tuple[bytes, List[Dict[str, Any]], int, int]:
    if not pptx:
        return file_bytes, [], 0, 0

    prs = pptx.Presentation(io.BytesIO(file_bytes))
    findings = []
    detected_count = 0
    redacted_count = 0

    def sanitize_text(text: str) -> str:
        nonlocal detected_count, redacted_count
        entities = analyze_text(text)
        if not entities:
            return text
        sanitized = text
        for e in entities:
            detected_count += 1
            findings.append({
                "entity_type": e["entity_type"],
                "confidence": e["score"],
                "severity": e["severity"],
                "action": e["action"],
                "occurrence_count": 1
            })
            matched = e.get("matched")
            if matched and matched in sanitized:
                sanitized = sanitized.replace(matched, f"<REDACTED_{e['entity_type']}>")
                redacted_count += 1
        return sanitized

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            run.text = sanitize_text(run.text)

    out_bytes = io.BytesIO()
    prs.save(out_bytes)
    return out_bytes.getvalue(), findings, detected_count, redacted_count
