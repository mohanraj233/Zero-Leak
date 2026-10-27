import io
from typing import Tuple, List, Dict, Any
from app.detectors.presidio_detector import analyze_text

try:
    import openpyxl
except ImportError:
    openpyxl = None

def redact_xlsx(file_bytes: bytes, filename: str) -> Tuple[bytes, List[Dict[str, Any]], int, int]:
    if not openpyxl:
        return file_bytes, [], 0, 0

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    findings = []
    detected_count = 0
    redacted_count = 0

    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    val = cell.value.strip()
                    entities = analyze_text(val)
                    if entities:
                        sanitized = val
                        for e in entities:
                            detected_count += 1
                            findings.append({
                                "entity_type": e["entity_type"],
                                "confidence": e["score"],
                                "severity": e["severity"],
                                "action": e["action"],
                                "occurrence_count": 1
                            })
                            matched = e.get("matched")
                            if matched and matched in sanitized:
                                sanitized = sanitized.replace(matched, f"<REDACTED_{e['entity_type']}>")
                                redacted_count += 1
                        cell.value = sanitized

    out_bytes = io.BytesIO()
    wb.save(out_bytes)
    return out_bytes.getvalue(), findings, detected_count, redacted_count
