import io
from typing import Tuple, List, Dict, Any
from app.detectors.presidio_detector import analyze_text

try:
    import docx
except ImportError:
    docx = None

def redact_docx(file_bytes: bytes, filename: str) -> Tuple[bytes, List[Dict[str, Any]], int, int]:
    if not docx:
        return file_bytes, [], 0, 0

    doc = docx.Document(io.BytesIO(file_bytes))
    findings = []
    detected_count = 0
    redacted_count = 0

    def sanitize_string(text: str) -> Tuple[str, int]:
        nonlocal detected_count, redacted_count
        entities = analyze_text(text)
        if not entities:
            return text, 0

        entities.sort(key=lambda x: x["start"], reverse=True)
        sanitized = text
        local_redacted = 0
        for e in entities:
            detected_count += 1
            findings.append({
                "entity_type": e["entity_type"],
                "confidence": e["score"],
                "severity": e["severity"],
                "action": e["action"],
                "occurrence_count": 1
            })
            matched = e.get("matched")
            if matched and matched in sanitized:
                sanitized = sanitized.replace(matched, f"<REDACTED_{e['entity_type']}>")
                local_redacted += 1
                redacted_count += 1
        return sanitized, local_redacted

    # 1. Inspect Paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            new_text, changed = sanitize_string(p.text)
            if changed > 0:
                p.text = new_text

    # 2. Inspect Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    new_text, changed = sanitize_string(cell.text)
                    if changed > 0:
                        cell.text = new_text

    out_bytes = io.BytesIO()
    doc.save(out_bytes)
    return out_bytes.getvalue(), findings, detected_count, redacted_count
